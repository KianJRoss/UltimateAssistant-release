from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from html.parser import HTMLParser
from pathlib import PurePath


MAX_FILE_BYTES = 15 * 1024 * 1024
MAX_EXTRACTED_CHARS = 30000
SUPPORTED_EXTENSIONS = {
    ".txt", ".md", ".csv", ".tsv", ".json", ".html", ".htm", ".xml",
    ".pdf", ".docx", ".pptx", ".xlsx", ".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp",
}


class _HTMLText(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._ignored = 0

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag.casefold() in {"script", "style", "noscript", "svg"}:
            self._ignored += 1

    def handle_endtag(self, tag: str) -> None:
        if tag.casefold() in {"script", "style", "noscript", "svg"} and self._ignored:
            self._ignored -= 1

    def handle_data(self, data: str) -> None:
        if not self._ignored and data.strip():
            self.parts.append(data.strip())


def _plain_text(raw: bytes) -> str:
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace")


def _check_office_archive(raw: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        unpacked = sum(entry.file_size for entry in entries)
        if len(entries) > 2000 or unpacked > 120 * 1024 * 1024:
            raise ValueError("This Office file expands beyond the safe processing limits.")
        if any(entry.file_size > 1024 * 1024 and entry.file_size / max(entry.compress_size, 1) > 150 for entry in entries):
            raise ValueError("This Office file has suspicious compression and was rejected.")


def _extract_pdf(raw: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw), strict=False)
    if reader.is_encrypted:
        raise ValueError("Encrypted PDFs are not supported; provide an unlocked copy.")
    if len(reader.pages) > 100:
        raise ValueError("PDFs are limited to 100 pages per attachment.")
    pages = []
    remaining = MAX_EXTRACTED_CHARS
    for number, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if text.strip():
            chunk = text[:remaining]
            pages.append(f"[Page {number}]\n{chunk}")
            remaining -= len(chunk)
        if remaining <= 0:
            break
    return "\n\n".join(pages)


def _extract_docx(raw: bytes) -> str:
    _check_office_archive(raw)
    from docx import Document

    document = Document(io.BytesIO(raw))
    pieces = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table_number, table in enumerate(document.tables, 1):
        rows = [" | ".join(cell.text.replace("\n", " ") for cell in row.cells) for row in table.rows]
        if rows:
            pieces.append(f"[Table {table_number}]\n" + "\n".join(rows))
    return "\n\n".join(pieces)


def _extract_pptx(raw: bytes) -> str:
    _check_office_archive(raw)
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(raw))
    if len(presentation.slides) > 200:
        raise ValueError("PowerPoint files are limited to 200 slides per attachment.")
    slides = []
    for number, slide in enumerate(presentation.slides, 1):
        text = [shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()]
        if text:
            slides.append(f"[Slide {number}]\n" + "\n".join(text))
    return "\n\n".join(slides)


def _extract_xlsx(raw: bytes) -> str:
    _check_office_archive(raw)
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheets = []
    try:
        if len(workbook.worksheets) > 40:
            raise ValueError("Excel files are limited to 40 worksheets per attachment.")
        for sheet in workbook.worksheets:
            rows = []
            for row_number, values in enumerate(sheet.iter_rows(values_only=True), 1):
                if row_number > 5000:
                    rows.append("[Row limit reached]")
                    break
                cells = ["" if value is None else str(value)[:1000] for value in values[:100]]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                sheets.append(f"[Worksheet: {sheet.title}]\n" + "\n".join(rows))
    finally:
        workbook.close()
    return "\n\n".join(sheets)


def extract_document(filename: str, raw: bytes) -> dict[str, object]:
    safe_name = PurePath(filename.replace("\\", "/")).name.strip()
    suffix = PurePath(safe_name).suffix.casefold()
    if not safe_name or suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}.")
    if not raw or len(raw) > MAX_FILE_BYTES:
        raise ValueError("Files must contain data and be no larger than 15 MB.")

    analysis_summary = ""
    if suffix in {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}:
        try:
            from herald.router.ingestion import ingest_bytes
        except ImportError as exc:
            raise ValueError("Herald's image ingestion is unavailable in this installation.") from exc
        result = ingest_bytes(raw, filename_hint=safe_name)
        analysis_summary = result.summary.strip()
        extracted = result.flat_text().strip()
        text = (f"[Image analysis: {analysis_summary}]\n" if analysis_summary else "") + extracted
    elif suffix in {".txt", ".md", ".csv", ".tsv", ".json", ".xml"}:
        text = _plain_text(raw)
        if suffix == ".json":
            try:
                text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
            except json.JSONDecodeError as exc:
                raise ValueError("The JSON file is malformed.") from exc
        elif suffix in {".csv", ".tsv"}:
            delimiter = "\t" if suffix == ".tsv" else ","
            text = "\n".join(" | ".join(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter))
    elif suffix in {".html", ".htm"}:
        parser = _HTMLText()
        parser.feed(_plain_text(raw))
        text = "\n".join(parser.parts)
    elif suffix == ".pdf":
        text = _extract_pdf(raw)
    elif suffix == ".docx":
        text = _extract_docx(raw)
    elif suffix == ".pptx":
        text = _extract_pptx(raw)
    elif suffix == ".xlsx":
        text = _extract_xlsx(raw)
    else:
        text = ""

    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise ValueError("No readable text was extracted. Scanned PDFs or images may need OCR/VLM dependencies.")
    return {
        "filename": safe_name,
        "text": text[:MAX_EXTRACTED_CHARS],
        "truncated": len(text) > MAX_EXTRACTED_CHARS,
        "format": suffix.lstrip("."),
        "summary": analysis_summary,
    }
